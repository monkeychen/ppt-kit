import io
import tempfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Literal

import pptx
import pymupdf
from PIL import Image, UnidentifiedImageError
from pptx.enum.shapes import MSO_SHAPE_TYPE

from pptkit.common.units import page_size_to_emu
from pptkit.render.renderer import render_pptx

SUPPORTED_IMAGE_EXTS: set[str] = {".png", ".jpg", ".jpeg", ".webp"}
SUPPORTED_EXTS: set[str] = SUPPORTED_IMAGE_EXTS | {".pdf", ".pptx"}


class IngestError(Exception):
    def __init__(self, code: str, message: str, hint: str) -> None:
        super().__init__(message)
        self.code = code
        self.message = message
        self.hint = hint


@dataclass
class IngestedPage:
    page_index: int
    route: Literal["vector", "raster"]
    width_px: int
    height_px: int
    width_emu: int
    height_emu: int
    raster_png_bytes: bytes


@dataclass
class IngestedDocument:
    source_path: Path
    pages: list[IngestedPage]


def _detect_pdf_page_route(page: pymupdf.Page) -> Literal["vector", "raster"]:
    raw_text = str(page.get_text("text")).strip()
    non_space_chars = [ch for ch in raw_text if not ch.isspace()]
    if len(non_space_chars) >= 5:
        valid_chars = sum(1 for ch in non_space_chars if ch != "\ufffd" and ch.isprintable())
        if valid_chars / len(non_space_chars) >= 0.9:
            return "vector"

    drawings = page.get_drawings()
    if len(drawings) >= 5:
        page_rect = page.rect
        page_area = max(page_rect.width * page_rect.height, 1.0)
        has_fullpage_raster = False
        for img_info in page.get_images(full=True):
            xref = img_info[0]
            for bbox in page.get_image_rects(xref):
                if (bbox.width * bbox.height) / page_area >= 0.9:
                    has_fullpage_raster = True
                    break
        if not has_fullpage_raster:
            return "vector"

    return "raster"


def _load_image_file(
    path: Path, route_override: Literal["auto", "vector", "raster"]
) -> IngestedDocument:
    try:
        with Image.open(path) as img:
            img.load()
            rgb = img.convert("RGB")
            width_px, height_px = rgb.size
            buf = io.BytesIO()
            rgb.save(buf, format="PNG")
            png_bytes = buf.getvalue()
    except (UnidentifiedImageError, OSError, ValueError) as exc:
        raise IngestError(
            code="CORRUPTED_FILE",
            message=f"无法读取图像文件：{path}",
            hint="文件已损坏或无法读取，请检查源文件完整性",
        ) from exc

    width_emu, height_emu = page_size_to_emu(width_px, height_px)
    route: Literal["vector", "raster"] = "raster" if route_override == "auto" else route_override
    return IngestedDocument(
        source_path=path,
        pages=[
            IngestedPage(
                page_index=0,
                route=route,
                width_px=width_px,
                height_px=height_px,
                width_emu=width_emu,
                height_emu=height_emu,
                raster_png_bytes=png_bytes,
            )
        ],
    )


def _load_pdf_file(
    path: Path, route_override: Literal["auto", "vector", "raster"]
) -> IngestedDocument:
    try:
        doc = pymupdf.open(path)
    except Exception as exc:
        raise IngestError(
            code="CORRUPTED_FILE",
            message=f"无法打开 PDF 文件：{path}",
            hint="文件已损坏或无法读取，请检查源文件完整性",
        ) from exc

    try:
        if doc.is_encrypted:
            raise IngestError(
                code="ENCRYPTED_PDF",
                message=f"PDF 文件已加密：{path}",
                hint=f"该 PDF 已加密，请先用 qpdf --decrypt {path.name} decrypted.pdf 解密",
            )
        if len(doc) == 0:
            raise IngestError(
                code="CORRUPTED_FILE",
                message=f"PDF 文件不包含任何页面：{path}",
                hint="请检查 PDF 文件是否为空或损坏",
            )

        pages: list[IngestedPage] = []
        for idx in range(len(doc)):
            page = doc[idx]
            route: Literal["vector", "raster"] = (
                _detect_pdf_page_route(page) if route_override == "auto" else route_override
            )
            pix = page.get_pixmap(dpi=144, alpha=False)
            width_px, height_px = pix.width, pix.height
            width_emu, height_emu = page_size_to_emu(width_px, height_px)
            pages.append(
                IngestedPage(
                    page_index=idx,
                    route=route,
                    width_px=width_px,
                    height_px=height_px,
                    width_emu=width_emu,
                    height_emu=height_emu,
                    raster_png_bytes=pix.tobytes("png"),
                )
            )
        return IngestedDocument(source_path=path, pages=pages)
    finally:
        doc.close()


def _shape_has_nonempty_text(shape: Any) -> bool:
    if not getattr(shape, "has_text_frame", False):
        return False
    tf = getattr(shape, "text_frame", None)
    if tf is None:
        return False
    return bool(getattr(tf, "text", "").strip())


def _extract_slide_image_bytes(slide: Any) -> bytes | None:
    for shape in slide.shapes:
        if _shape_has_nonempty_text(shape):
            return None

    pictures = [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
    if len(pictures) == 1 and len(slide.shapes) == 1:
        try:
            return pictures[0].image.blob
        except Exception:
            return None

    if len(slide.shapes) == 0:
        try:
            image_rels = [
                rel.target_part
                for rel in slide.part.rels.values()
                if hasattr(rel, "reltype") and "image" in rel.reltype
            ]
            if len(image_rels) == 1 and hasattr(image_rels[0], "blob"):
                return image_rels[0].blob
        except Exception:
            return None

    return None


def _load_pptx_file(
    path: Path, route_override: Literal["auto", "vector", "raster"]
) -> IngestedDocument:
    try:
        prs = pptx.Presentation(str(path))
    except Exception as exc:
        raise IngestError(
            code="CORRUPTED_FILE",
            message=f"无法打开 PPTX 文件：{path}",
            hint="文件已损坏或不是合法的 PPTX 文件，请检查源文件完整性",
        ) from exc

    if len(prs.slides) == 0:
        raise IngestError(
            code="CORRUPTED_FILE",
            message=f"PPTX 文件不包含任何页面：{path}",
            hint="请检查 PPTX 文件是否为空或损坏",
        )

    slide_w_emu: int = int(getattr(prs, "slide_width", 0) or 12192000)
    slide_h_emu: int = int(getattr(prs, "slide_height", 0) or 6858000)

    extracted_blobs: list[bytes | None] = [
        _extract_slide_image_bytes(slide) for slide in prs.slides
    ]
    all_direct = all(blob is not None for blob in extracted_blobs)

    pages: list[IngestedPage] = []
    if all_direct:
        for idx, blob in enumerate(extracted_blobs):
            assert blob is not None
            try:
                with Image.open(io.BytesIO(blob)) as img:
                    rgb = img.convert("RGB")
                    width_px, height_px = rgb.size
                    buf = io.BytesIO()
                    rgb.save(buf, format="PNG")
                    png_bytes = buf.getvalue()
            except Exception as exc:
                raise IngestError(
                    code="CORRUPTED_FILE",
                    message=f"无法解析 PPTX 第 {idx + 1} 页的图片数据",
                    hint="该幻灯片中的图片格式损坏，请检查源文件",
                ) from exc

            route: Literal["vector", "raster"] = (
                "raster" if route_override == "auto" else route_override
            )
            pages.append(
                IngestedPage(
                    page_index=idx,
                    route=route,
                    width_px=width_px,
                    height_px=height_px,
                    width_emu=slide_w_emu,
                    height_emu=slide_h_emu,
                    raster_png_bytes=png_bytes,
                )
            )
    else:
        try:
            with tempfile.TemporaryDirectory() as tmp_dir_str:
                tmp_dir = Path(tmp_dir_str)
                render_res = render_pptx(path, tmp_dir)
                for idx, png_p in enumerate(render_res.png_paths):
                    route_detected: Literal["vector", "raster"] = "raster"
                    route = route_detected if route_override == "auto" else route_override

                    with Image.open(png_p) as img:
                        rgb = img.convert("RGB")
                        width_px, height_px = rgb.size
                        buf = io.BytesIO()
                        rgb.save(buf, format="PNG")
                        png_bytes = buf.getvalue()

                    pages.append(
                        IngestedPage(
                            page_index=idx,
                            route=route,
                            width_px=width_px,
                            height_px=height_px,
                            width_emu=slide_w_emu,
                            height_emu=slide_h_emu,
                            raster_png_bytes=png_bytes,
                        )
                    )
        except IngestError:
            raise
        except Exception as exc:
            raise IngestError(
                code="RENDER_FAILED",
                message=f"无法渲染 PPTX 页面：{exc}",
                hint="需要 Microsoft PowerPoint 或 LibreOffice 进行幻灯片无头渲染",
            ) from exc

    return IngestedDocument(source_path=path, pages=pages)


def load_input(
    path: Path | str, route_override: Literal["auto", "vector", "raster"] = "auto"
) -> IngestedDocument:
    file_path = Path(path)
    if not file_path.exists() or not file_path.is_file():
        raise IngestError(
            code="FILE_NOT_FOUND",
            message=f"输入文件不存在：{file_path}",
            hint="请确认输入文件路径是否正确",
        )

    ext = file_path.suffix.lower()
    if ext not in SUPPORTED_EXTS:
        raise IngestError(
            code="UNSUPPORTED_FORMAT",
            message=f"不支持的文件格式：{ext or '<无扩展名>'}",
            hint="支持的格式：.png, .jpg, .jpeg, .webp, .pdf, .pptx",
        )

    if ext == ".pdf":
        return _load_pdf_file(file_path, route_override)
    if ext == ".pptx":
        return _load_pptx_file(file_path, route_override)
    return _load_image_file(file_path, route_override)
