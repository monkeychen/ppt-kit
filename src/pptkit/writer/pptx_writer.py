from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Emu, Pt

from pptkit.common.color import hex_to_rgb
from pptkit.ir.models import (
    Background,
    CellBorder,
    ImageElement,
    PresentationIR,
    ShapeElement,
    SlideElement,
    SlideIR,
    SlideMaster,
    TableElement,
    TextBoxElement,
)

ALIGN_MAP = {
    "left": PP_ALIGN.LEFT,
    "center": PP_ALIGN.CENTER,
    "right": PP_ALIGN.RIGHT,
    "justify": PP_ALIGN.JUSTIFY,
}

VERTICAL_ANCHOR_MAP = {
    "top": MSO_ANCHOR.TOP,
    "middle": MSO_ANCHOR.MIDDLE,
    "bottom": MSO_ANCHOR.BOTTOM,
}

SHAPE_TYPE_MAP = {
    "rect": MSO_SHAPE.RECTANGLE,
    "round_rect": MSO_SHAPE.ROUNDED_RECTANGLE,
    "ellipse": MSO_SHAPE.OVAL,
    "right_arrow": MSO_SHAPE.RIGHT_ARROW,
}


def _rgb_color(hex_str: str) -> RGBColor:
    r, g, b = hex_to_rgb(hex_str)
    return RGBColor(r, g, b)


def _set_east_asian_font(run: Any, font_family: str) -> None:
    rpr = run._r.get_or_add_rPr()
    for existing in rpr.findall("{http://schemas.openxmlformats.org/drawingml/2006/main}ea"):
        rpr.remove(existing)
    ea = OxmlElement("a:ea")
    ea.set("typeface", font_family)
    rpr.append(ea)


def _apply_background(slide: Any, bg: Background, width_emu: int, height_emu: int) -> None:
    if bg.fill_type == "image" and bg.image_path:
        pic = slide.shapes.add_picture(
            bg.image_path, Emu(0), Emu(0), Emu(width_emu), Emu(height_emu)
        )
        pic.name = "bg-image"
    else:
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = _rgb_color(bg.color_hex)


def _add_textbox(slide: Any, elem: TextBoxElement) -> None:
    box = slide.shapes.add_textbox(
        Emu(elem.bbox.x),
        Emu(elem.bbox.y),
        Emu(elem.bbox.width),
        Emu(elem.bbox.height),
    )
    box.name = elem.id
    tf = box.text_frame
    tf.word_wrap = elem.word_wrap
    tf.margin_left = Emu(elem.margin_left_emu)
    tf.margin_right = Emu(elem.margin_right_emu)
    tf.margin_top = Emu(elem.margin_top_emu)
    tf.margin_bottom = Emu(elem.margin_bottom_emu)

    for p_idx, para_ir in enumerate(elem.paragraphs):
        p = tf.paragraphs[0] if p_idx == 0 else tf.add_paragraph()
        p.alignment = ALIGN_MAP[para_ir.align]
        if para_ir.line_spacing_pt is not None:
            p.line_spacing = Pt(para_ir.line_spacing_pt)
        if para_ir.space_after_pt > 0:
            p.space_after = Pt(para_ir.space_after_pt)

        for run_ir in para_ir.runs:
            run = p.add_run()
            run.text = run_ir.text
            font = run.font
            font.name = run_ir.style.font_family
            font.size = Pt(run_ir.style.font_size_pt)
            font.bold = run_ir.style.bold
            font.italic = run_ir.style.italic
            font.underline = run_ir.style.underline
            font.color.rgb = _rgb_color(run_ir.style.color_hex)
            _set_east_asian_font(run, run_ir.style.font_family)


def _disable_shadow(shape: Any) -> None:
    try:
        shape.shadow.inherit = False
    except Exception:
        pass
    a_ns = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    p_ns = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
    for style_el in shape._element.findall(f"{p_ns}style"):
        for eff_ref in style_el.findall(f"{a_ns}effectRef"):
            eff_ref.set("idx", "0")


def _apply_fill_and_stroke(shape: Any, elem: ShapeElement) -> None:
    _disable_shadow(shape)
    if elem.fill_color_hex:
        shape.fill.solid()
        shape.fill.fore_color.rgb = _rgb_color(elem.fill_color_hex)
    else:
        shape.fill.background()

    if elem.stroke_color_hex and elem.stroke_width_emu > 0:
        shape.line.color.rgb = _rgb_color(elem.stroke_color_hex)
        shape.line.width = Emu(elem.stroke_width_emu)
    else:
        shape.line.fill.background()


def _add_shape(slide: Any, elem: ShapeElement) -> None:
    b = elem.bbox
    if elem.shape_type == "line":
        pts = elem.points_emu or [(b.x, b.y), (b.x + b.width, b.y + b.height)]
        conn = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT,
            Emu(pts[0][0]),
            Emu(pts[0][1]),
            Emu(pts[-1][0]),
            Emu(pts[-1][1]),
        )
        conn.name = elem.id
        _disable_shadow(conn)
        if elem.stroke_color_hex:
            conn.line.color.rgb = _rgb_color(elem.stroke_color_hex)
        conn.line.width = Emu(max(elem.stroke_width_emu, 1))
        return

    if elem.shape_type == "freeform" and elem.points_emu and len(elem.points_emu) >= 3:
        x0, y0 = elem.points_emu[0]
        builder = slide.shapes.build_freeform(Emu(x0), Emu(y0))
        builder.add_line_segments(
            [(Emu(px), Emu(py)) for px, py in elem.points_emu[1:]], close=True
        )
        shape = builder.convert_to_shape()
        shape.name = elem.id
        _apply_fill_and_stroke(shape, elem)
        return

    mso_type = SHAPE_TYPE_MAP.get(elem.shape_type, MSO_SHAPE.RECTANGLE)
    shape = slide.shapes.add_shape(mso_type, Emu(b.x), Emu(b.y), Emu(b.width), Emu(b.height))
    shape.name = elem.id
    if elem.shape_type == "round_rect" and len(shape.adjustments) > 0:
        shape.adjustments[0] = elem.corner_radius_ratio
    _apply_fill_and_stroke(shape, elem)


def _add_image(slide: Any, elem: ImageElement) -> None:
    b = elem.bbox
    pic = slide.shapes.add_picture(elem.image_path, Emu(b.x), Emu(b.y), Emu(b.width), Emu(b.height))
    pic.name = elem.id


def _apply_cell_border(cell: Any, side_tag: str, border: CellBorder | None) -> None:
    if not border or border.width_emu <= 0:
        return
    tcPr = cell._tc.get_or_add_tcPr()
    ns = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    for existing in tcPr.findall(f"{ns}{side_tag}"):
        tcPr.remove(existing)
    ln = OxmlElement(f"a:{side_tag}")
    ln.set("w", str(border.width_emu))
    solidFill = OxmlElement("a:solidFill")
    srgbClr = OxmlElement("a:srgbClr")
    srgbClr.set("val", border.color_hex.lstrip("#").upper())
    solidFill.append(srgbClr)
    ln.append(solidFill)
    tcPr.append(ln)


def _add_table(slide: Any, elem: TableElement) -> None:
    b = elem.bbox
    shape = slide.shapes.add_table(
        elem.rows,
        elem.cols,
        Emu(b.x),
        Emu(b.y),
        Emu(b.width),
        Emu(b.height),
    )
    shape.name = elem.id
    _disable_shadow(shape)
    table = shape.table

    # Set column widths
    for col_idx, w_emu in enumerate(elem.col_widths_emu):
        if col_idx < len(table.columns):
            table.columns[col_idx].width = Emu(w_emu)

    # Set row heights
    for row_idx, h_emu in enumerate(elem.row_heights_emu):
        if row_idx < len(table.rows):
            table.rows[row_idx].height = Emu(h_emu)

    # Populate cells
    for r_idx in range(min(elem.rows, len(elem.cells))):
        row_cells = elem.cells[r_idx]
        for c_idx in range(min(elem.cols, len(row_cells))):
            cell_ir = row_cells[c_idx]
            cell = table.cell(r_idx, c_idx)

            # Zero margins or use IR values
            cell.margin_left = Emu(cell_ir.margin_left_emu)
            cell.margin_right = Emu(cell_ir.margin_right_emu)
            cell.margin_top = Emu(cell_ir.margin_top_emu)
            cell.margin_bottom = Emu(cell_ir.margin_bottom_emu)

            # Vertical alignment
            cell.vertical_anchor = VERTICAL_ANCHOR_MAP.get(
                cell_ir.vertical_align, MSO_ANCHOR.MIDDLE
            )

            # Fill
            if cell_ir.fill_color_hex:
                cell.fill.solid()
                cell.fill.fore_color.rgb = _rgb_color(cell_ir.fill_color_hex)
            else:
                cell.fill.background()

            # Borders
            _apply_cell_border(cell, "lnL", cell_ir.border_left)
            _apply_cell_border(cell, "lnR", cell_ir.border_right)
            _apply_cell_border(cell, "lnT", cell_ir.border_top)
            _apply_cell_border(cell, "lnB", cell_ir.border_bottom)

            # Paragraphs and text runs
            tf = cell.text_frame
            tf.word_wrap = True
            if cell_ir.paragraphs:
                for p_idx, para_ir in enumerate(cell_ir.paragraphs):
                    p = tf.paragraphs[0] if p_idx == 0 else tf.add_paragraph()
                    p.alignment = ALIGN_MAP.get(para_ir.align, PP_ALIGN.LEFT)
                    if para_ir.line_spacing_pt is not None:
                        p.line_spacing = Pt(para_ir.line_spacing_pt)
                    if para_ir.space_after_pt > 0:
                        p.space_after = Pt(para_ir.space_after_pt)

                    for run_ir in para_ir.runs:
                        run = p.add_run()
                        run.text = run_ir.text
                        font = run.font
                        font.name = run_ir.style.font_family
                        font.size = Pt(run_ir.style.font_size_pt)
                        font.bold = run_ir.style.bold
                        font.italic = run_ir.style.italic
                        font.underline = run_ir.style.underline
                        font.color.rgb = _rgb_color(run_ir.style.color_hex)
                        _set_east_asian_font(run, run_ir.style.font_family)
            else:
                tf.text = ""

    # Merge spanned cells
    for r_idx in range(min(elem.rows, len(elem.cells))):
        row_cells = elem.cells[r_idx]
        for c_idx in range(min(elem.cols, len(row_cells))):
            cell_ir = row_cells[c_idx]
            if cell_ir.row_span > 1 or cell_ir.col_span > 1:
                target_r = min(elem.rows - 1, r_idx + cell_ir.row_span - 1)
                target_c = min(elem.cols - 1, c_idx + cell_ir.col_span - 1)
                table.cell(r_idx, c_idx).merge(table.cell(target_r, target_c))


def _write_slide(
    pres: Any,
    slide_ir: SlideIR,
    master_map: dict[str, SlideMaster] | None = None,
) -> None:
    blank_layout = pres.slide_layouts[6]
    slide = pres.slides.add_slide(blank_layout)

    master = master_map.get(slide_ir.master_id) if (master_map and slide_ir.master_id) else None
    bg = slide_ir.background
    if master and (
        (bg.fill_type == "solid" and bg.color_hex == "#FFFFFF" and not bg.image_path)
        or (bg.fill_type == "image" and not bg.image_path)
    ):
        if master.background.fill_type == "image" and master.background.image_path:
            bg = master.background

    _apply_background(slide, bg, slide_ir.width_emu, slide_ir.height_emu)

    elements_to_render: list[SlideElement] = []
    if master and master.static_elements:
        elements_to_render.extend(master.static_elements)
    elements_to_render.extend(slide_ir.elements)

    sorted_elements = sorted(elements_to_render, key=lambda e: e.z_order)
    for elem in sorted_elements:
        if isinstance(elem, TextBoxElement):
            _add_textbox(slide, elem)
        elif isinstance(elem, ShapeElement):
            _add_shape(slide, elem)
        elif isinstance(elem, ImageElement):
            _add_image(slide, elem)
        elif isinstance(elem, TableElement):
            _add_table(slide, elem)


def write_pptx(ir: PresentationIR, output_path: Path | str) -> Path:
    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)

    pres = Presentation()
    pres.slide_width = Emu(ir.width_emu)
    pres.slide_height = Emu(ir.height_emu)

    master_map = {m.id: m for m in ir.masters}

    for slide_ir in ir.slides:
        _write_slide(pres, slide_ir, master_map)

    pres.save(str(out))
    return out
