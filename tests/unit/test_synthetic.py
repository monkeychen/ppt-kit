from pathlib import Path
from typing import Any, cast

import pptx

from benchmark.synthetic.evaluator import evaluate_slide_ir
from benchmark.synthetic.generator import (
    generate_synthetic_dataset,
    generate_synthetic_table_dataset,
)
from pptkit.ir.models import BBox, ShapeElement, TextBoxElement


def test_generate_synthetic_slides_and_evaluate_metrics(tmp_path: Path) -> None:
    items = generate_synthetic_dataset(out_dir=tmp_path, count=2, seed=42, render=False)
    assert len(items) == 2
    for item in items:
        assert item.pptx_path.exists()
        assert item.gt_json_path.exists()

    gt_slide = items[0].slide_ir

    # 1. Perfect prediction equals ground truth -> perfect scores
    perfect = evaluate_slide_ir(gt_slide, gt_slide, ssim=1.0)
    assert perfect.char_accuracy == 1.0
    assert perfect.textbox_iou == 1.0
    assert perfect.font_size_mae_pt == 0.0
    assert perfect.color_delta_e == 0.0
    assert perfect.shape_recall == 1.0
    assert perfect.shape_precision == 1.0
    assert perfect.ssim == 1.0

    # 2. Perturbed prediction -> measurable error
    pred_slide = gt_slide.model_copy(deep=True)
    for elem in pred_slide.elements:
        if isinstance(elem, TextBoxElement):
            elem.bbox = BBox(
                x=elem.bbox.x + 100000,
                y=elem.bbox.y,
                width=elem.bbox.width,
                height=elem.bbox.height,
            )
            elem.paragraphs[0].runs[0].style.font_size_pt += 2.0
            elem.paragraphs[0].runs[0].style.color_hex = "#334455"
        elif isinstance(elem, ShapeElement):
            elem.shape_type = "ellipse"

    perturbed = evaluate_slide_ir(pred_slide, gt_slide, ssim=0.95)
    assert 0.5 < perturbed.textbox_iou < 1.0
    assert perturbed.font_size_mae_pt > 0.0
    assert perturbed.color_delta_e > 0.0
    assert perturbed.ssim == 0.95


def test_generate_synthetic_table_slides(tmp_path: Path) -> None:
    items = generate_synthetic_table_dataset(out_dir=tmp_path, count=2, seed=123, render=False)
    assert len(items) == 2
    for item in items:
        assert item.pptx_path.exists()
        assert item.gt_json_path.exists()

        prs = pptx.Presentation(str(item.pptx_path))
        shapes = prs.slides[0].shapes
        table_shapes = [s for s in shapes if s.has_table]
        assert len(table_shapes) == 1
        table = cast(Any, table_shapes[0]).table
        assert len(table.rows) == 4
        assert len(table.columns) == 4
        assert "模块名称" in table.cell(0, 0).text
        assert "vector.parser" in table.cell(1, 0).text
