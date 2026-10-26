from pathlib import Path
from typing import Any, cast

import pptx
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

from pptkit.ir.models import (
    BBox,
    CellBorder,
    Paragraph,
    PresentationIR,
    SlideIR,
    TableCell,
    TableElement,
    TextRun,
    TextStyle,
)
from pptkit.writer.pptx_writer import write_pptx


def test_write_pptx_creates_native_table_shape(tmp_path: Path) -> None:
    c00 = TableCell(
        row=0,
        col=0,
        paragraphs=[
            Paragraph(
                runs=[
                    TextRun(
                        text="标题1",
                        style=TextStyle(font_size_pt=14.0, color_hex="#111827", bold=True),
                    )
                ],
                align="center",
            )
        ],
        fill_color_hex="#E5E7EB",
        border_bottom=CellBorder(color_hex="#374151", width_emu=19050),
    )
    c01 = TableCell(
        row=0,
        col=1,
        paragraphs=[
            Paragraph(
                runs=[
                    TextRun(
                        text="标题2",
                        style=TextStyle(font_size_pt=14.0, color_hex="#111827", bold=True),
                    )
                ],
                align="center",
            )
        ],
        fill_color_hex="#E5E7EB",
        border_bottom=CellBorder(color_hex="#374151", width_emu=19050),
    )
    c10 = TableCell(
        row=1,
        col=0,
        paragraphs=[
            Paragraph(
                runs=[
                    TextRun(
                        text="数值A",
                        style=TextStyle(font_size_pt=12.0, color_hex="#4B5563"),
                    )
                ],
                align="left",
            )
        ],
        fill_color_hex="#F9FAFB",
    )
    c11 = TableCell(
        row=1,
        col=1,
        paragraphs=[
            Paragraph(
                runs=[
                    TextRun(
                        text="数值B",
                        style=TextStyle(font_size_pt=12.0, color_hex="#4B5563"),
                    )
                ],
                align="right",
            )
        ],
        fill_color_hex="#FFFFFF",
    )

    elem = TableElement(
        id="table-main",
        bbox=BBox(x=1000000, y=1500000, width=7000000, height=3000000),
        z_order=0,
        rows=2,
        cols=2,
        col_widths_emu=[3500000, 3500000],
        row_heights_emu=[1000000, 2000000],
        cells=[[c00, c01], [c10, c11]],
    )

    slide_ir = SlideIR(page_index=0, route="vector", elements=[elem])
    pres_ir = PresentationIR(source_path="test_src", slides=[slide_ir])

    out_file = tmp_path / "table_output.pptx"
    write_pptx(pres_ir, out_file)

    assert out_file.exists()
    prs = pptx.Presentation(str(out_file))
    slide = prs.slides[0]
    assert len(slide.shapes) == 1
    shape = slide.shapes[0]
    assert shape.has_table is True
    table = cast(Any, shape).table

    assert len(table.rows) == 2
    assert len(table.columns) == 2
    assert table.columns[0].width == Emu(3500000)
    assert table.columns[1].width == Emu(3500000)

    # Check cell 0,0
    cell_00 = table.cell(0, 0)
    assert cell_00.text.strip() == "标题1"
    p00 = cell_00.text_frame.paragraphs[0]
    assert p00.alignment == PP_ALIGN.CENTER
    assert p00.runs[0].font.bold is True
    assert p00.runs[0].font.size == Pt(14.0)

    # Margins should be zeroed (or set to IR values)
    assert cell_00.margin_left == 0
    assert cell_00.margin_right == 0
    assert cell_00.margin_top == 0
    assert cell_00.margin_bottom == 0
    assert cell_00.vertical_anchor == MSO_ANCHOR.MIDDLE

    # DrawingML Border check on cell 0,0
    tcPr = cell_00._tc.get_or_add_tcPr()
    lnB = tcPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}lnB")
    assert lnB is not None
    assert lnB.get("w") == "19050"


def test_table_cell_spans_merge(tmp_path: Path) -> None:
    # 2x2 table where row 0 spans 2 columns
    c00 = TableCell(
        row=0,
        col=0,
        col_span=2,
        paragraphs=[
            Paragraph(
                runs=[
                    TextRun(
                        text="合并表头",
                        style=TextStyle(font_size_pt=16.0, color_hex="#000000"),
                    )
                ]
            )
        ],
    )
    c01 = TableCell(row=0, col=1)  # covered by span
    c10 = TableCell(
        row=1,
        col=0,
        paragraphs=[
            Paragraph(
                runs=[
                    TextRun(
                        text="左下",
                        style=TextStyle(font_size_pt=12.0, color_hex="#000000"),
                    )
                ]
            )
        ],
    )
    c11 = TableCell(
        row=1,
        col=1,
        paragraphs=[
            Paragraph(
                runs=[TextRun(text="右下", style=TextStyle(font_size_pt=12.0, color_hex="#000000"))]
            )
        ],
    )

    elem = TableElement(
        id="table-merged",
        bbox=BBox(x=500000, y=500000, width=6000000, height=2000000),
        z_order=0,
        rows=2,
        cols=2,
        col_widths_emu=[3000000, 3000000],
        row_heights_emu=[1000000, 1000000],
        cells=[[c00, c01], [c10, c11]],
    )

    slide_ir = SlideIR(page_index=0, route="vector", elements=[elem])
    pres_ir = PresentationIR(source_path="test_merged", slides=[slide_ir])

    out_file = tmp_path / "merged_table.pptx"
    write_pptx(pres_ir, out_file)

    prs = pptx.Presentation(str(out_file))
    table = cast(Any, prs.slides[0].shapes[0]).table
    merged_cell = table.cell(0, 0)
    assert merged_cell.is_merge_origin is True
    assert merged_cell.span_width == 2
