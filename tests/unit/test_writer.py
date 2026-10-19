from pathlib import Path
from typing import Any

from PIL import Image
from pptx import Presentation

from pptkit.common.units import page_size_to_emu, pt_to_emu
from pptkit.ir.models import (
    Background,
    BBox,
    ImageElement,
    Paragraph,
    PresentationIR,
    ShapeElement,
    SlideIR,
    TextBoxElement,
    TextRun,
    TextStyle,
)
from pptkit.writer.pptx_writer import write_pptx


def test_write_pptx_native_textbox_and_shapes_editable(tmp_path: Path) -> None:
    # Create a tiny test image
    img_path = tmp_path / "crop.png"
    Image.new("RGB", (100, 80), color=(20, 120, 220)).save(img_path)

    # Custom 4:3 aspect ratio
    w_emu, h_emu = page_size_to_emu(1600, 1200)

    # Purposely place elements out of z_order in list to verify writer sorts by z_order (AC3.4)
    tb = TextBoxElement(
        id="tb-title",
        bbox=BBox(x=500000, y=400000, width=8000000, height=1200000),
        z_order=30,
        paragraphs=[
            Paragraph(
                runs=[
                    TextRun(
                        text="可编辑中英文标题 Editable Title",
                        style=TextStyle(
                            font_family="PingFang SC",
                            font_size_pt=28.0,
                            color_hex="#112233",
                            bold=True,
                        ),
                    )
                ],
                align="center",
                line_spacing_pt=34.0,
            )
        ],
    )
    rect = ShapeElement(
        id="sh-card",
        bbox=BBox(x=400000, y=300000, width=8500000, height=2000000),
        z_order=10,
        shape_type="round_rect",
        fill_color_hex="#EBF3FA",
        stroke_color_hex="#2B6CB0",
        stroke_width_emu=pt_to_emu(2.0),
        corner_radius_ratio=0.1,
    )
    freeform = ShapeElement(
        id="sh-poly",
        bbox=BBox(x=1000000, y=3000000, width=2000000, height=1500000),
        z_order=20,
        shape_type="freeform",
        fill_color_hex="#F6AD55",
        points_emu=[(1000000, 3000000), (3000000, 3000000), (2000000, 4500000)],
    )
    pic = ImageElement(
        id="im-photo",
        bbox=BBox(x=5000000, y=3000000, width=2500000, height=2000000),
        z_order=5,
        image_path=str(img_path),
    )

    slide = SlideIR(
        page_index=0,
        route="vector",
        width_emu=w_emu,
        height_emu=h_emu,
        background=Background(fill_type="solid", color_hex="#FFFFFF"),
        elements=[tb, rect, freeform, pic],
    )
    pres_ir = PresentationIR(
        source_path="synthetic.pdf",
        width_emu=w_emu,
        height_emu=h_emu,
        slides=[slide],
    )

    out_pptx = tmp_path / "out.pptx"
    write_pptx(pres_ir, out_pptx)

    # Re-open with python-pptx and verify AC3.2, AC3.3, AC3.4
    loaded = Presentation(str(out_pptx))
    assert loaded.slide_width == w_emu
    assert loaded.slide_height == h_emu
    assert len(loaded.slides) == 1

    shapes: list[object] = list(loaded.slides[0].shapes)
    assert len(shapes) == 4

    # Verify z_order sorting: pic (5) -> rect (10) -> freeform (20) -> tb (30)
    assert [getattr(s, "name", "") for s in shapes] == [
        "im-photo",
        "sh-card",
        "sh-poly",
        "tb-title",
    ]

    # Verify native shape is editable and has expected fill & geometry
    card_shape: Any = shapes[1]
    assert card_shape.left == 400000
    assert card_shape.width == 8500000
    assert str(card_shape.fill.fore_color.rgb) == "EBF3FA"

    # Verify native textbox is editable and East Asian font (<a:ea>) is written
    title_shape: Any = shapes[3]
    assert title_shape.has_text_frame
    assert title_shape.text_frame.text == "可编辑中英文标题 Editable Title"
    rpr_xml = title_shape.text_frame.paragraphs[0].runs[0]._r.xml
    assert 'typeface="PingFang SC"' in rpr_xml
    assert "a:ea" in rpr_xml
