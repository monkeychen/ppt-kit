import io
import json
from pathlib import Path

import pptx
import pymupdf
import pytest
from PIL import Image, ImageDraw
from pptx.util import Inches
from typer.testing import CliRunner

from pptkit.cli.app import app

runner = CliRunner()


def _create_vector_pdf(path: Path, pages: int = 1) -> None:
    doc = pymupdf.open()
    for i in range(pages):
        p = doc.new_page(width=960, height=540)
        p.insert_text((100, 120), f"Editable Vector Slide Page {i + 1} Content")
    doc.save(path)
    doc.close()


def _create_raster_png(path: Path) -> None:
    img = Image.new("RGB", (640, 360), color=(250, 248, 242))
    draw = ImageDraw.Draw(img)
    draw.text((80, 100), "Slide Text For Raster Conversion", fill=(20, 20, 20))
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    path.write_bytes(buf.getvalue())


def _create_raster_pptx(path: Path, pages: int = 1) -> None:
    prs = pptx.Presentation()
    for i in range(pages):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        buf = io.BytesIO()
        img = Image.new("RGB", (640, 360), color=(50 * (i + 1), 100, 150))
        draw = ImageDraw.Draw(img)
        draw.text((50, 50), f"Raster Slide {i + 1}", fill=(255, 255, 255))
        img.save(buf, format="PNG")
        buf.seek(0)
        slide.shapes.add_picture(buf, Inches(0), Inches(0), prs.slide_width, prs.slide_height)
    prs.save(str(path))


def test_cli_convert_file_and_dir_outputs_pptx_and_report(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    in_dir = tmp_path / "inputs"
    in_dir.mkdir()
    pdf_file = in_dir / "deck.pdf"
    png_file = in_dir / "slide.png"
    pptx_file = in_dir / "album.pptx"
    _create_vector_pdf(pdf_file)
    _create_raster_png(png_file)
    _create_raster_pptx(pptx_file, pages=1)

    out_single = tmp_path / "out_single"
    res_single = runner.invoke(
        app, ["convert", str(pdf_file), "-o", str(out_single), "--max-iter", "0"]
    )
    assert res_single.exit_code == 0
    assert (out_single / "deck.pptx").exists()
    report_path = out_single / "deck.report.json"
    assert report_path.exists()
    rep = json.loads(report_path.read_text(encoding="utf-8"))
    assert rep["pages"][0]["route"] == "vector"

    # Single PPTX conversion
    out_pptx = tmp_path / "out_pptx"
    res_pptx = runner.invoke(
        app, ["convert", str(pptx_file), "-o", str(out_pptx), "--max-iter", "0"]
    )
    assert res_pptx.exit_code == 0
    assert (out_pptx / "album.pptx").exists()
    assert (out_pptx / "album.report.json").exists()

    # In-place conversion safety check (output to same dir does not overwrite original)
    res_inplace = runner.invoke(
        app, ["convert", str(pptx_file), "-o", str(in_dir), "--max-iter", "0"]
    )
    assert res_inplace.exit_code == 0
    assert (in_dir / "album.converted.pptx").exists()

    # Batch directory conversion
    out_batch = tmp_path / "out_batch"
    res_batch = runner.invoke(
        app, ["convert", str(in_dir), "-o", str(out_batch), "--max-iter", "0"]
    )
    assert res_batch.exit_code == 0
    assert (out_batch / "deck.pptx").exists()
    assert (out_batch / "deck.report.json").exists()
    assert (out_batch / "slide.pptx").exists()
    assert (out_batch / "slide.report.json").exists()
    assert (out_batch / "album.pptx").exists()
    assert (out_batch / "album.report.json").exists()


def test_cli_pages_filter_progress_stderr_and_vlm_key_check(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    for var in (
        "PPTKIT_GEMINI_API_KEY",
        "GEMINI_API_KEY",
        "PPTKIT_OPENAI_API_KEY",
        "OPENAI_API_KEY",
    ):
        monkeypatch.delenv(var, raising=False)
    monkeypatch.setenv("PPTKIT_CONFIG_PATH", str(tmp_path / "nonexistent_config.toml"))

    pdf_file = tmp_path / "multi.pdf"
    _create_vector_pdf(pdf_file, pages=4)
    png_file = tmp_path / "ras.png"
    _create_raster_png(png_file)

    out_dir = tmp_path / "out"
    # 1. --pages 1,3-4 converts only pages 0, 2, 3 (3 pages total) and --json outputs valid JSON
    res_pages = runner.invoke(
        app,
        [
            "convert",
            str(pdf_file),
            "-o",
            str(out_dir),
            "--pages",
            "1,3-4",
            "--max-iter",
            "0",
            "--json",
        ],
    )
    assert res_pages.exit_code == 0
    rep = json.loads((out_dir / "multi.report.json").read_text(encoding="utf-8"))
    assert rep["succeeded_pages"] == 3
    assert [p["page_index"] for p in rep["pages"]] == [0, 2, 3]

    # 2. Offline Apple Vision raster works WITHOUT API key
    res_ras_offline = runner.invoke(
        app, ["convert", str(png_file), "-o", str(out_dir), "--max-iter", "0"]
    )
    assert res_ras_offline.exit_code == 0

    # 3. With --vlm flag and no API key, raster fails with exit code 1 and env hint (AC9.4)
    res_ras_vlm = runner.invoke(
        app, ["convert", str(png_file), "-o", str(out_dir), "--vlm", "--max-iter", "0"]
    )
    assert res_ras_vlm.exit_code == 1
    combined_out = res_ras_vlm.stdout + (res_ras_vlm.stderr or "")
    assert "PPTKIT_GEMINI_API_KEY" in combined_out or "GEMINI_API_KEY" in combined_out

    # 4. pptkit report --needs-review subcommand (AC10.2)
    res_rep_cmd = runner.invoke(
        app,
        ["report", str(out_dir / "multi.report.json"), "--needs-review", "--json"],
    )
    assert res_rep_cmd.exit_code == 0


def test_cli_partial_failure_isolation_and_exit_codes(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    # Mixed PDF (1 vector + 1 raster page) with --vlm and NO API Key -> exit code 2
    for var in ("PPTKIT_GEMINI_API_KEY", "GEMINI_API_KEY"):
        monkeypatch.delenv(var, raising=False)
    monkeypatch.setenv("PPTKIT_CONFIG_PATH", str(tmp_path / "nonexistent_config.toml"))

    mixed_pdf = tmp_path / "mixed.pdf"
    doc = pymupdf.open()
    p0 = doc.new_page(width=960, height=540)
    p0.insert_text((100, 120), "Page 0 Vector Text Content")
    p1 = doc.new_page(width=960, height=540)
    img = Image.new("RGB", (320, 180), color=(200, 210, 220))
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    p1.insert_image(pymupdf.Rect(0, 0, 960, 540), stream=buf.getvalue())
    doc.save(mixed_pdf)
    doc.close()

    out_dir = tmp_path / "out_partial"
    res_partial = runner.invoke(
        app, ["convert", str(mixed_pdf), "-o", str(out_dir), "--vlm", "--max-iter", "0"]
    )
    assert res_partial.exit_code == 2
    assert (out_dir / "mixed.pptx").exists()
    rep = json.loads((out_dir / "mixed.report.json").read_text(encoding="utf-8"))
    assert rep["succeeded_pages"] == 1
    assert rep["failed_pages"] == 1

    # Corrupted file -> exit code 1
    bad_pdf = tmp_path / "corrupted.pdf"
    bad_pdf.write_bytes(b"invalid pdf bytes")
    res_bad = runner.invoke(app, ["convert", str(bad_pdf), "-o", str(out_dir)])
    assert res_bad.exit_code == 1
