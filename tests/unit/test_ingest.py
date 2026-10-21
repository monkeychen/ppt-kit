import io
from pathlib import Path

import pptx
import pymupdf
import pytest
from PIL import Image, ImageDraw
from pptx.util import Inches

from pptkit.ingest.loader import IngestError, load_input


def _make_raster_image_bytes(fmt: str, size: tuple[int, int] = (320, 180)) -> bytes:
    img = Image.new("RGB", size, color=(240, 244, 248))
    draw = ImageDraw.Draw(img)
    draw.rectangle([20, 20, 120, 80], fill=(30, 100, 200))
    buf = io.BytesIO()
    img.save(buf, format=fmt)
    return buf.getvalue()


def _create_vector_pdf(path: Path, pages: int = 1) -> None:
    doc = pymupdf.open()
    for i in range(pages):
        page = doc.new_page(width=960, height=540)
        page.insert_text((80, 100), f"Slide {i + 1}: Vector Editable Title Text")
        page.draw_rect(pymupdf.Rect(80, 150, 300, 260), color=(0, 0, 1), fill=(0.8, 0.9, 1.0))
    doc.save(path)
    doc.close()


def _create_raster_pdf(path: Path, pages: int = 1) -> None:
    doc = pymupdf.open()
    png_bytes = _make_raster_image_bytes("PNG", (640, 360))
    for _ in range(pages):
        page = doc.new_page(width=960, height=540)
        page.insert_image(pymupdf.Rect(0, 0, 960, 540), stream=png_bytes)
    doc.save(path)
    doc.close()


def _create_raster_pptx(path: Path, pages: int = 1) -> None:
    prs = pptx.Presentation()
    for i in range(pages):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        buf = io.BytesIO()
        img = Image.new("RGB", (640, 360), color=(50 * (i + 1), 100, 150))
        draw = ImageDraw.Draw(img)
        draw.text((50, 50), f"Raster Slide {i + 1}", fill=(255, 255, 255))
        img.save(buf, format="PNG")
        buf.seek(0)
        slide.shapes.add_picture(buf, Inches(0), Inches(0), prs.slide_width, prs.slide_height)
    prs.save(str(path))


def test_load_six_formats_page_count_matches(tmp_path: Path) -> None:
    formats = [
        ("sample.png", "PNG"),
        ("sample.jpg", "JPEG"),
        ("sample.jpeg", "JPEG"),
        ("sample.webp", "WEBP"),
    ]
    for filename, pil_fmt in formats:
        fpath = tmp_path / filename
        fpath.write_bytes(_make_raster_image_bytes(pil_fmt, (640, 360)))
        doc = load_input(fpath)
        assert len(doc.pages) == 1
        assert doc.pages[0].width_px == 640
        assert doc.pages[0].height_px == 360
        assert doc.pages[0].route == "raster"

    pdf_path = tmp_path / "sample.pdf"
    _create_vector_pdf(pdf_path, pages=3)
    pdf_doc = load_input(pdf_path)
    assert len(pdf_doc.pages) == 3

    pptx_path = tmp_path / "sample.pptx"
    _create_raster_pptx(pptx_path, pages=2)
    pptx_doc = load_input(pptx_path)
    assert len(pptx_doc.pages) == 2
    assert pptx_doc.pages[0].route == "raster"
    assert pptx_doc.pages[1].route == "raster"
    assert len(pptx_doc.pages[0].raster_png_bytes) > 0


def test_route_detection_vector_vs_raster_and_mixed_pdf(tmp_path: Path) -> None:
    mixed_pdf = tmp_path / "mixed.pdf"
    doc = pymupdf.open()

    # Page 0: vector text + shape
    p0 = doc.new_page(width=960, height=540)
    p0.insert_text((100, 120), "Editable Vector Presentation Slide")
    p0.draw_rect(pymupdf.Rect(100, 160, 400, 300), color=(0, 0, 0))

    # Page 1: full-page raster image (like NotebookLM export)
    p1 = doc.new_page(width=960, height=540)
    p1.insert_image(pymupdf.Rect(0, 0, 960, 540), stream=_make_raster_image_bytes("PNG"))

    doc.save(mixed_pdf)
    doc.close()

    loaded = load_input(mixed_pdf)
    assert [p.route for p in loaded.pages] == ["vector", "raster"]


def test_route_benchmark_20_annotated_samples_100_accuracy(tmp_path: Path) -> None:
    samples: list[tuple[Path, str]] = []

    # 10 vector samples (PPT-exported PDF style: 8 with text+shapes, 2 with pure vector shapes)
    for i in range(8):
        p = tmp_path / f"pptpdf_{i:02d}.pdf"
        _create_vector_pdf(p, pages=1)
        samples.append((p, "vector"))

    for i in range(8, 10):
        p = tmp_path / f"pptpdf_shapes_{i:02d}.pdf"
        doc = pymupdf.open()
        page = doc.new_page(width=960, height=540)
        for k in range(6):
            page.draw_rect(
                pymupdf.Rect(50 + k * 120, 100, 140 + k * 120, 220),
                color=(0.1, 0.2, 0.5),
                fill=(0.7, 0.8, 0.9),
            )
        doc.save(p)
        doc.close()
        samples.append((p, "vector"))

    # 5 NotebookLM-style raster PDFs
    for i in range(5):
        p = tmp_path / f"notebooklm_{i:02d}.pdf"
        _create_raster_pdf(p, pages=1)
        samples.append((p, "raster"))

    # 5 single-image slides (PNG, JPG, WEBP)
    exts = [("png", "PNG"), ("jpg", "JPEG"), ("jpeg", "JPEG"), ("webp", "WEBP"), ("png", "PNG")]
    for i, (ext, fmt) in enumerate(exts):
        p = tmp_path / f"ai_slide_{i:02d}.{ext}"
        p.write_bytes(_make_raster_image_bytes(fmt))
        samples.append((p, "raster"))

    assert len(samples) == 20
    correct = 0
    for sample_path, expected_route in samples:
        loaded = load_input(sample_path)
        if loaded.pages[0].route == expected_route:
            correct += 1
    assert correct == 20


def test_ingest_error_on_encrypted_corrupted_unsupported(tmp_path: Path) -> None:
    # 1. Unsupported format
    txt_file = tmp_path / "notes.txt"
    txt_file.write_text("not a slide")
    with pytest.raises(IngestError) as exc_info:
        load_input(txt_file)
    assert exc_info.value.code == "UNSUPPORTED_FORMAT"
    assert ".pdf" in exc_info.value.hint

    # 2. Encrypted PDF
    enc_pdf = tmp_path / "encrypted.pdf"
    doc = pymupdf.open()
    page = doc.new_page(width=960, height=540)
    page.insert_text((100, 100), "Secret slide")
    doc.save(
        enc_pdf,
        encryption=getattr(pymupdf, "PDF_ENCRYPT_AES_256", 4),
        owner_pw="owner",
        user_pw="secret",
    )
    doc.close()

    with pytest.raises(IngestError) as exc_info:
        load_input(enc_pdf)
    assert exc_info.value.code == "ENCRYPTED_PDF"
    assert "qpdf --decrypt" in exc_info.value.hint

    # 3. Corrupted file
    bad_pdf = tmp_path / "corrupted.pdf"
    bad_pdf.write_bytes(b"not a valid pdf header")
    with pytest.raises(IngestError) as exc_info:
        load_input(bad_pdf)
    assert exc_info.value.code == "CORRUPTED_FILE"

    bad_png = tmp_path / "corrupted.png"
    bad_png.write_bytes(b"not a valid png")
    with pytest.raises(IngestError) as exc_info:
        load_input(bad_png)
    assert exc_info.value.code == "CORRUPTED_FILE"

    # 4. Corrupted PPTX
    bad_pptx = tmp_path / "corrupted.pptx"
    bad_pptx.write_bytes(b"not a valid zip or pptx")
    with pytest.raises(IngestError) as exc_info:
        load_input(bad_pptx)
    assert exc_info.value.code == "CORRUPTED_FILE"

    # 5. Empty PPTX (0 slides)
    empty_pptx = tmp_path / "empty.pptx"
    pptx.Presentation().save(str(empty_pptx))
    with pytest.raises(IngestError) as exc_info:
        load_input(empty_pptx)
    assert exc_info.value.code == "CORRUPTED_FILE"
