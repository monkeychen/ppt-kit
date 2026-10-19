from pathlib import Path
from typing import Any

from pptx import Presentation

from pptkit.cli.app import _convert_single_file
from pptkit.render.renderer import render_with_powerpoint


def test_convert_user_real_slide_to_editable_pptx(tmp_path: Path) -> None:
    real_img = (
        Path(__file__).resolve().parents[2] / "benchmark/real/other/01-ai-beginner-workflow.jpg"
    )
    assert real_img.exists()

    out_dir = Path("out/real_m2")
    succ, fail, rep = _convert_single_file(
        src_file=real_img,
        out_dir=out_dir,
        route_opt="raster",
        provider="gemini",
        max_iter=1,
        use_vlm=False,
    )
    assert succ == 1 and fail == 0
    pptx_path = out_dir / "01-ai-beginner-workflow.pptx"
    assert pptx_path.exists()
    assert rep.pages[0].editable_elements >= 75
    assert rep.pages[0].ssim is not None and rep.pages[0].ssim > 0.65
    assert rep.pages[0].text_region_mae is not None and rep.pages[0].text_region_mae < 60.0

    # Open generated PPTX and verify native editable textboxes, geometric shapes, and image layers
    pres = Presentation(str(pptx_path))
    shapes: list[Any] = list(pres.slides[0].shapes)
    tb_shapes = [s for s in shapes if s.name.startswith("p0-tb-")]
    geom_shapes = [s for s in shapes if s.name.startswith("p0-sh-")]
    img_layers = [s for s in shapes if s.name.startswith("p0-img-")]
    assert len(tb_shapes) >= 35
    assert len(geom_shapes) >= 20
    assert len(img_layers) >= 10

    all_texts = [s.text_frame.text for s in shapes if getattr(s, "has_text_frame", False)]
    joined = "\n".join(all_texts)

    assert "如果你想用好AI你就按着一步一步做" in joined
    assert "一个非常非常初级的小流程" in joined
    assert "需要一台电脑" in joined
    assert "背景、痛点、需求" in joined
    assert "请你帮我直接完成吧" in joined
    assert "同时沉淀成Skill" in joined
    assert "现在，只需要半天时间，你已经学会使用AI了。" in joined

    # Also render the final PPTX via PowerPoint to out/real_m2/rendered/ for visual inspection
    render_with_powerpoint(pptx_path, out_dir / "rendered", target_size=(2068, 1162))
    assert (out_dir / "rendered" / "01-ai-beginner-workflow_page_000.png").exists()
